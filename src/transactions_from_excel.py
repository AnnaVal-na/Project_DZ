from typing import Any, Dict, List
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
import logging
from datetime import datetime

logging.basicConfig(level=logging.ERROR)


def load_transactions_from_excel(file_path: str) -> List[Dict[str, Any]]:
    transactions = []
    workbook = None
    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        if not workbook.sheetnames:
            raise ValueError("Файл не содержит листов.")

        sheet: Worksheet = workbook.active  # type: ignore
        if sheet is None:
            raise ValueError("Активный лист не найден")

        # Получаем заголовки
        header_row = sheet[1]
        headers = [str(cell.value).strip() if cell.value else "" for cell in header_row]

        # Обработка данных
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) != len(headers):
                continue

            transaction: Dict[str, Any] = {}
            for i, key in enumerate(headers):
                value = row[i]
                if key == "Amount":
                    if isinstance(value, (int, float)):
                        transaction[key] = int(value)
                    elif isinstance(value, str):
                        try:
                            transaction[key] = int(float(value))
                        except ValueError:
                            transaction[key] = 0
                    elif isinstance(value, datetime):
                        transaction[key] = value.toordinal()
                    else:
                        transaction[key] = 0
                else:
                    transaction[key] = value
            transactions.append(transaction)

    except FileNotFoundError:  # Убрано 'as e'
        logging.error(f"Файл не найден: {file_path}")
        raise
    except Exception as e:
        logging.error(f"Ошибка при чтении файла: {e}")
        raise
    finally:
        if workbook:
            workbook.close()
    return transactions
